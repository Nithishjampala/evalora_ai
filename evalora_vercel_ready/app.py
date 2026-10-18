import io, json, os, re
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import fitz
import gspread
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from google import genai
from google.oauth2.service_account import Credentials
from pptx import Presentation

app = FastAPI(title="Evalora AI", version="1.0.0")
MAX_BYTES = 8 * 1024 * 1024
ALLOWED_PHASES = {"Ideation", "Design", "Build", "Demo"}
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md", ".csv"}
DEFAULT_MODELS = ["gemini-3.8-flash",
    "gemini-3.6-flash",
    "gemini-3.5-flash-lite"]
MODELS = [os.getenv(f"GEMINI_MODEL_{i}", DEFAULT_MODELS[i-1]).strip() for i in range(1, 4)]


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Supported files: PDF, DOCX, PPTX, TXT, MD, CSV.")
    if ext == ".pdf":
        with fitz.open(stream=data, filetype="pdf") as pdf:
            return "\n".join(page.get_text() for page in pdf)
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        chunks = [p.text for p in doc.paragraphs]
        chunks.extend(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        return "\n".join(chunks)
    if ext == ".pptx":
        prs = Presentation(io.BytesIO(data))
        return "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    return data.decode("utf-8-sig", errors="replace")


def generate_with_fallback(prompt: str) -> dict:
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise RuntimeError("Missing GEMINI_API_KEY. Add it in Vercel → Settings → Environment Variables.")
    client = genai.Client(api_key=api_key)
    errors = []
    seen = set()
    for model in MODELS:
        if not model or model in seen:
            continue
        seen.add(model)
        try:
            response = client.models.generate_content(model=model, contents=prompt)
            answer = (getattr(response, "text", None) or "").strip()
            if answer:
                return {"text": answer, "model_used": model}
            errors.append(f"{model}: empty response")
        except Exception as exc:
            errors.append(f"{model}: {str(exc)[:160]}")
    raise RuntimeError("All configured Gemini models failed. Check model availability, API quota, and API key. " + " | ".join(errors))


def parse_json(text: str) -> dict:
    cleaned = re.sub(r"^\s*```(?:json)?\s*|\s*```\s*$", "", text.strip(), flags=re.I)
    try:
        result = json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\{[\s\S]*\}", cleaned)
        if not match:
            raise ValueError("Gemini returned a response that was not valid JSON. Please retry.")
        result = json.loads(match.group(0))
    if not isinstance(result, dict):
        raise ValueError("Gemini returned an unexpected JSON format. Please retry.")
    return result


def save_to_sheets(row: dict) -> str:
    service_json = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON")
    sheet_id = os.getenv("GOOGLE_SHEET_ID")
    if not service_json or not sheet_id:
        return "Not saved: configure GOOGLE_SERVICE_ACCOUNT_JSON and GOOGLE_SHEET_ID in Vercel."
    info = json.loads(service_json)
    creds = Credentials.from_service_account_info(info, scopes=["https://www.googleapis.com/auth/spreadsheets"])
    book = gspread.authorize(creds).open_by_key(sheet_id)
    worksheet_name = os.getenv("GOOGLE_WORKSHEET", "Evaluations")
    try:
        ws = book.worksheet(worksheet_name)
    except gspread.WorksheetNotFound:
        ws = book.add_worksheet(title=worksheet_name, rows=1000, cols=9)
    headers = ["Timestamp", "Name", "Project ID", "Project Name", "Phase", "Score", "Report", "Guideline Compliance", "Status"]
    if not ws.row_values(1):
        ws.append_row(headers, value_input_option="RAW")
    ws.append_row([
        datetime.now(timezone.utc).isoformat(), row["name"], row["project_id"], row["project_name"],
        row["phase"], row["score"], row["report"], row["compliance"], "Completed"
    ], value_input_option="RAW")
    return "Saved to Google Sheets."


@app.get("/api/health")
def health():
    return {"ok": True, "configured_models": MODELS, "gemini_key_configured": bool(os.getenv("GEMINI_API_KEY"))}


@app.post("/api/evaluate")
async def evaluate(
    name: str = Form(...),
    project_id: str = Form(...),
    project_name: str = Form(...),
    phase: str = Form(...),
    guideline_text: str = Form(""),
    guideline_file: Optional[UploadFile] = File(None),
    submission_file: UploadFile = File(...),
):
    if phase not in ALLOWED_PHASES:
        raise HTTPException(400, "Invalid phase. Choose Ideation, Design, Build, or Demo.")
    try:
        sub_data = await submission_file.read()
        if len(sub_data) > MAX_BYTES:
            raise HTTPException(413, "Submission file exceeds 8 MB.")
        submission = extract_text(submission_file.filename, sub_data)
        if guideline_file and guideline_file.filename:
            gd = await guideline_file.read()
            if len(gd) > MAX_BYTES:
                raise HTTPException(413, "Guideline file exceeds 8 MB.")
            guideline_text += "\n" + extract_text(guideline_file.filename, gd)
        if not submission.strip():
            raise HTTPException(400, "No readable text found. If this is a scanned PDF, OCR it first or upload a text-based document.")
        if not guideline_text.strip():
            guideline_text = "Assess phase-appropriate innovation, problem-solution fit, feasibility, impact, validation, clarity and completeness. Clearly distinguish evidence from assumptions."

        project_context = f"Project: {project_name}; ID: {project_id}; phase: {phase}\nGUIDELINES:\n{guideline_text[:12000]}\nSUBMISSION:\n{submission[:20000]}"
        p1 = "Act as a careful project evaluator. Use only evidence in the submission. Return ONLY valid JSON with strengths (array), weaknesses (array), missing_evidence (array), guideline_compliance (string), provisional_score (number 0-100), rationale (string).\n" + project_context
        r1 = generate_with_fallback(p1)
        p2 = "Independently audit the first evaluation. Do not assume unsupported claims are true. Return ONLY valid JSON with score (number 0-100), strengths (array), weaknesses (array), guideline_compliance (string), rationale (string).\n" + project_context + "\nFIRST REVIEW:\n" + r1["text"][:7000]
        r2 = generate_with_fallback(p2)
        p3 = "Reconcile the two reviews conservatively. Do not recommend a shortlist. Return ONLY valid JSON with score (number 0-100), guideline_compliance (string), report (Markdown string), strengths (array), improvements (array), missing_evidence (array), score_explanation (string). The report must include project name, project ID, phase, score, evidence-based strengths, improvements, missing evidence, and score explanation.\n" + project_context + "\nREVIEW 1:\n" + r1["text"][:5000] + "\nREVIEW 2:\n" + r2["text"][:5000]
        r3 = generate_with_fallback(p3)
        final = parse_json(r3["text"])
        try:
            score = max(0.0, min(100.0, float(final.get("score", 0))))
        except (TypeError, ValueError):
            score = 0.0
        report = str(final.get("report") or f"# {project_name}\n\n**Score:** {score}/100\n\n{final.get('score_explanation', '')}")
        compliance = str(final.get("guideline_compliance", "Not determined"))
        row = {"name": name.strip(), "project_id": project_id.strip(), "project_name": project_name.strip(), "phase": phase, "score": score, "report": report, "compliance": compliance}
        try:
            sheet_status = save_to_sheets(row)
        except Exception as exc:
            sheet_status = "Google Sheets save failed: " + str(exc)[:220]
        return {"success": True, "score": score, "report": report, "guideline_compliance": compliance, "sheet_status": sheet_status, "models_used": [r1["model_used"], r2["model_used"], r3["model_used"]]}
    except HTTPException:
        raise
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        # Keep diagnostic detail useful without returning credentials.
        raise HTTPException(500, str(exc)[:700]) from exc
