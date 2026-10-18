
import os
import json
import re
import traceback
from pathlib import Path
from http.server import BaseHTTPRequestHandler
from email.parser import BytesParser
from email.policy import default

from dotenv import load_dotenv
import fitz
import docx
from pptx import Presentation
from striprtf.striprtf import rtf_to_text
from google import genai
from google.genai import types
import gspread
from google.oauth2.service_account import Credentials


ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env.local")

MODELS = [
    "gemini-3.8-flash",
    "gemini-3.6-flash",
    "gemini-3.5-flash-lite",
]

SHEET_ID = os.getenv(
    "GOOGLE_SHEET_ID",
    "1L8FuY_CFURGrRkWbEFAG_QsHTyF0oUJYfpCfNAl-sOw"
)
WORKSHEET_NAME = os.getenv("GOOGLE_WORKSHEET_NAME", "Evaluations")

HEADERS = [
    "Name",
    "Project ID",
    "Project Name",
    "Phase Type",
    "Questions",
    "AI Confidence Score",
    "Feedback",
    "Suggestions",
    "Overall Score",
]

RATING_POINTS = {
    "Outstanding": 100,
    "Very Good": 85,
    "Good": 70,
    "Needs Improvement": 40,
    "Incomplete": 0,
}

COMPLIANCE_STATUSES = {
    "satisfied": "Satisfied",
    "partially satisfied": "Partially Satisfied",
    "not satisfied": "Not Satisfied",
    "not applicable": "Not Applicable",
}


def extract_document(filename, content):
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        with fitz.open(stream=content, filetype="pdf") as pdf:
            return "\n".join(page.get_text() for page in pdf)

    if ext == ".docx":
        from io import BytesIO
        document = docx.Document(BytesIO(content))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    if ext == ".pptx":
        from io import BytesIO
        presentation = Presentation(BytesIO(content))
        parts = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts.append(f"Slide {slide_number}")
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    parts.append(shape.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    if ext == ".txt":
        return content.decode("utf-8", errors="replace")

    if ext == ".rtf":
        return rtf_to_text(content.decode("utf-8", errors="replace"))

    raise ValueError(
        "Unsupported file type. Upload PDF, DOCX, PPTX, TXT, or RTF."
    )


def parse_multipart(content_type, body):
    message = BytesParser(policy=default).parsebytes(
        b"Content-Type: " + content_type.encode("utf-8") +
        b"\r\nMIME-Version: 1.0\r\n\r\n" + body
    )

    if not message.is_multipart():
        raise ValueError("Invalid multipart form. Please try again.")

    fields = {}
    for part in message.iter_parts():
        name = part.get_param(
            "name", header="content-disposition"
        )
        if not name:
            continue

        filename = part.get_filename()
        payload = part.get_payload(decode=True) or b""

        if filename:
            fields[name] = {
                "filename": filename,
                "content": payload,
            }
        else:
            charset = part.get_content_charset() or "utf-8"
            fields[name] = payload.decode(charset, errors="replace")

    return fields


def get_text_file(fields, field_name, required=False):
    item = fields.get(field_name)

    if not isinstance(item, dict):
        if required:
            raise ValueError(f"Please upload {field_name}.")
        return ""

    if not item["content"]:
        if required:
            raise ValueError(f"The {field_name} file is empty.")
        return ""

    return extract_document(item["filename"], item["content"])


def get_client():
    key = os.getenv("GEMINI_API_KEY")
    if not key:
        raise ValueError("GEMINI_API_KEY is missing from .env.local.")
    return genai.Client(api_key=key)


def request_ai(client, prompt):
    errors = []

    for model in MODELS:
        try:
            response = client.models.generate_content(
                model=model,
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    temperature=0.2,
                ),
            )
            if not response.text:
                raise ValueError("Gemini returned an empty response.")
            return json.loads(response.text), model
        except Exception as exc:
            errors.append(f"{model}: {str(exc)[:350]}")

    raise RuntimeError(
        "All configured Gemini model attempts failed. Check model access, "
        "API key, and API availability. Details: " + " | ".join(errors)
    )


def build_prompt(name, project_id, project_name, phase,
                 submission_text, guideline_text):
    return f"""
You are Evalora AI, a careful project evaluation assistant.

Evaluate the provided submission for the selected phase: {phase}.
Do not invent evidence. If evidence is absent, state that clearly.
Distinguish between explicit evidence, reasonable inference, and missing
information. Do not assume a requirement is met just because the submission
uses similar wording. Compare the actual meaning and evidence.

PROJECT INFORMATION
Evaluator: {name}
Project ID: {project_id}
Project name: {project_name}
Selected phase: {phase}

PHASE EXPECTATIONS
Ideation: problem clarity, target users, need, novelty, feasibility,
impact, and validation of the idea.
Design: system design, architecture, workflow, components, implementation
plan, and alignment with requirements.
Build: implementation evidence, working prototype, technical decisions,
testing, reliability, and progress.
Demo: demonstrated functionality, results, validation, impact, limitations,
and clarity of demonstration.

SCORING
Q1-Q5 must each have exactly one rating:
Outstanding = 100
Very Good = 85
Good = 70
Needs Improvement = 40
Incomplete = 0

Q6 is additional feedback only. It is not scored and is not included in
the overall score. Evaluate the same core areas for all phases, but adapt
the expectations to the selected phase. Consider relevance, evidence,
technical quality, feasibility, and phase readiness.

GUIDELINE COMPLIANCE
Assess each important guideline separately using only these statuses:
Satisfied, Partially Satisfied, Not Satisfied, Not Applicable.
Explain the evidence or gap. Do not judge only by exact word matching.
Use Not Applicable only when a requirement genuinely does not apply.
If no guidelines are supplied, return an empty guideline_compliance array.

Return ONLY valid JSON in this schema:
{{
  "project_name": "...",
  "project_id": "...",
  "phase": "{phase}",
  "questions": [
    {{
      "id": "Q1",
      "question": "Problem and user need",
      "rating": "Outstanding",
      "score": 100,
      "feedback": "Evidence-based explanation"
    }},
    {{
      "id": "Q2",
      "question": "Innovation and relevance",
      "rating": "Good",
      "score": 70,
      "feedback": "Evidence-based explanation"
    }},
    {{
      "id": "Q3",
      "question": "Technical approach and feasibility",
      "rating": "Very Good",
      "score": 85,
      "feedback": "Evidence-based explanation"
    }},
    {{
      "id": "Q4",
      "question": "Phase deliverables and completeness",
      "rating": "Needs Improvement",
      "score": 40,
      "feedback": "Evidence-based explanation"
    }},
    {{
      "id": "Q5",
      "question": "Validation, impact, and readiness",
      "rating": "Incomplete",
      "score": 0,
      "feedback": "Evidence-based explanation"
    }},
    {{
      "id": "Q6",
      "question": "Additional feedback",
      "rating": "Not Scored",
      "score": null,
      "feedback": "Additional observations; not scored"
    }}
  ],
  "guideline_compliance": [
    {{
      "requirement": "Requirement text or concise description",
      "status": "Satisfied",
      "evidence": "Evidence or explanation"
    }}
  ],
  "ai_confidence_score": 75,
  "overall_feedback": "Balanced summary of strengths and gaps",
  "suggestions": [
    "Specific actionable improvement"
  ]
}}

Ensure Q1-Q5 use the exact rating names and scores shown above.
Ensure Q6 is never included in the score.
The overall score will be calculated separately as the average of Q1-Q5.

GUIDELINES
{guideline_text[:30000] if guideline_text else "No explicit guidelines supplied."}

PROJECT SUBMISSION
{submission_text[:50000]}
"""


def normalize_result(result, project_name, project_id, phase):
    result["project_name"] = project_name
    result["project_id"] = project_id
    result["phase"] = phase

    incoming = result.get("questions", [])
    by_id = {
        str(item.get("id", "")).upper(): item
        for item in incoming if isinstance(item, dict)
    }

    normalized_questions = []
    for qid, default_question in [
        ("Q1", "Problem and user need"),
        ("Q2", "Innovation and relevance"),
        ("Q3", "Technical approach and feasibility"),
        ("Q4", "Phase deliverables and completeness"),
        ("Q5", "Validation, impact, and readiness"),
        ("Q6", "Additional feedback"),
    ]:
        item = by_id.get(qid, {})
        rating = str(item.get("rating", "Incomplete")).strip()

        if qid == "Q6":
            normalized_questions.append({
                "id": "Q6",
                "question": item.get("question", default_question),
                "rating": "Not Scored",
                "score": None,
                "feedback": str(
                    item.get("feedback", "No additional feedback provided.")
                ),
            })
            continue

        if rating not in RATING_POINTS:
            rating = "Incomplete"

        normalized_questions.append({
            "id": qid,
            "question": str(item.get("question", default_question)),
            "rating": rating,
            "score": RATING_POINTS[rating],
            "feedback": str(item.get("feedback", "No explanation provided.")),
        })

    result["questions"] = normalized_questions
    result["overall_score"] = round(
        sum(q["score"] for q in normalized_questions if q["id"] != "Q6") / 5
    )

    confidence = result.get("ai_confidence_score", 0)
    try:
        confidence = int(float(confidence))
    except (ValueError, TypeError):
        confidence = 0
    result["ai_confidence_score"] = max(0, min(100, confidence))

    compliance = result.get("guideline_compliance", [])
    clean_compliance = []
    for item in compliance if isinstance(compliance, list) else []:
        if not isinstance(item, dict):
            continue
        raw_status = str(item.get("status", "Not Applicable")).strip().lower()
        status = COMPLIANCE_STATUSES.get(raw_status, "Not Applicable")
        clean_compliance.append({
            "requirement": str(item.get("requirement", "Requirement")),
            "status": status,
            "evidence": str(item.get("evidence", item.get("explanation", ""))),
        })
    result["guideline_compliance"] = clean_compliance

    if not isinstance(result.get("suggestions"), list):
        result["suggestions"] = []

    result["overall_feedback"] = str(
        result.get("overall_feedback", "No overall feedback provided.")
    )
    result["additional_feedback"] = normalized_questions[-1]["feedback"]
    return result


def save_to_sheets(evaluator, result):
    raw_credentials = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON")
    if not raw_credentials:
        raise ValueError("GOOGLE_SERVICE_ACCOUNT_JSON is missing from .env.local.")

    try:
        service_info = json.loads(raw_credentials)
    except json.JSONDecodeError as exc:
        raise ValueError(
            "GOOGLE_SERVICE_ACCOUNT_JSON is not valid JSON. Put the full "
            "service-account JSON on one line in .env.local."
        ) from exc

    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ]
    credentials = Credentials.from_service_account_info(
        service_info, scopes=scopes
    )
    client = gspread.authorize(credentials)
    spreadsheet = client.open_by_key(SHEET_ID)

    try:
        worksheet = spreadsheet.worksheet(WORKSHEET_NAME)
    except gspread.WorksheetNotFound:
        worksheet = spreadsheet.add_worksheet(
            title=WORKSHEET_NAME, rows=1000, cols=len(HEADERS)
        )
        worksheet.update(
            range_name="A1:I1",
            values=[HEADERS],
        )

    existing = worksheet.row_values(1)
    if not existing:
        worksheet.update(range_name="A1:I1", values=[HEADERS])
    elif existing[:len(HEADERS)] != HEADERS:
        raise ValueError(
            "The Evaluations sheet headers do not match the required nine "
            "columns. Update row 1 to the exact required headers."
        )

    scored_questions = [
        {
            "id": q["id"],
            "question": q["question"],
            "rating": q["rating"],
            "score": q["score"],
            "feedback": q["feedback"],
        }
        for q in result["questions"] if q["id"] != "Q6"
    ]

    row = [
        evaluator,
        result["project_id"],
        result["project_name"],
        result["phase"],
        json.dumps(scored_questions, ensure_ascii=False),
        result["ai_confidence_score"],
        result["overall_feedback"],
        json.dumps(result["suggestions"], ensure_ascii=False),
        result["overall_score"],
    ]

    if len(row) != 9:
        raise ValueError("Internal error: expected exactly nine spreadsheet values.")

    worksheet.append_row(row, value_input_option="USER_ENTERED")


class handler(BaseHTTPRequestHandler):
    def send_json(self, status, payload):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        self.send_json(200, {
            "ok": True,
            "message": "Evalora AI evaluation API is running."
        })

    def do_POST(self):
        if self.path.split("?")[0] != "/api/evaluate":
            self.send_json(404, {"error": "API route not found."})
            return

        try:
            length = int(self.headers.get("Content-Length", "0"))
            if length <= 0:
                raise ValueError("The request body is empty.")
            if length > 4 * 1024 * 1024:
                raise ValueError("Request exceeds the 4 MB limit.")

            content_type = self.headers.get("Content-Type", "")
            if "multipart/form-data" not in content_type:
                raise ValueError("Expected a multipart form submission.")

            fields = parse_multipart(
                content_type, self.rfile.read(length)
            )

            evaluator = str(fields.get("evaluator_name", "")).strip()
            project_id = str(fields.get("project_id", "")).strip()
            project_name = str(fields.get("project_name", "")).strip()
            phase = str(fields.get("phase", "")).strip()
            guidelines_text = str(fields.get("guidelines_text", "")).strip()

            if not evaluator or not project_id or not project_name:
                raise ValueError("Evaluator, project ID, and project name are required.")

            if phase not in {"Ideation", "Design", "Build", "Demo"}:
                raise ValueError("Select a valid evaluation phase.")

            submission_text = get_text_file(
                fields, "submission_file", required=True
            )
            uploaded_guidelines = get_text_file(
                fields, "guideline_file", required=False
            )

            guideline_text = "\n\n".join(
                text for text in [guidelines_text, uploaded_guidelines] if text
            )

            if not guideline_text:
                raise ValueError("Enter guidelines or upload a guideline file.")

            if not submission_text.strip():
                raise ValueError("No readable text was extracted from the submission.")

            client = get_client()
            prompt = build_prompt(
                evaluator, project_id, project_name, phase,
                submission_text, guideline_text
            )
            result, model_used = request_ai(client, prompt)
            result = normalize_result(result, project_name, project_id, phase)

            save_to_sheets(evaluator, result)

            self.send_json(200, {
                "ok": True,
                "model_used": model_used,
                "result": result,
            })

        except ValueError as exc:
            self.send_json(400, {"error": str(exc)})
        except Exception:
            traceback.print_exc()
            self.send_json(500, {
                "error": (
                    "Evaluation failed. Check the VS Code terminal for the "
                    "server-side details, including Gemini and Google Sheets access."
                )
            })
