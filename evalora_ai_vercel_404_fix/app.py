import os,re,json,io
from pathlib import Path
from fastapi import FastAPI,UploadFile,File,Form,HTTPException
import fitz,docx,gspread
from pptx import Presentation
from striprtf.striprtf import rtf_to_text
from google import genai
from google.genai import types
from google.oauth2.service_account import Credentials

app=FastAPI(title="Evalora AI")
SPREADSHEET_ID=os.getenv("SPREADSHEET_ID","1L8FuY_CFURGrRkWbEFAG_QsHTyF0oUJYfpCfNAl-sOw")
WORKSHEET=os.getenv("GOOGLE_WORKSHEET","Evaluations")
MODELS=[os.getenv("GEMINI_MODEL_2","gemini-3.6-flash"),os.getenv("GEMINI_MODEL_1","gemini-3.8-flash"),os.getenv("GEMINI_MODEL_3","gemini-3.5-flash-lite")]
RATING_SCORES={"Outstanding":100,"Very Good":85,"Good":70,"Needs Improvement":40,"Incomplete":0}
BASE_HEADERS=["Name","Project ID","Project Name","Phase Type","Questions","AI Confidence Score","Feedback","Suggestions","Overall Score"]
EXTRA_HEADERS=["Model Name","Answer Accuracy (%)"]
HEADERS=BASE_HEADERS+EXTRA_HEADERS
PHASES={
"Ideation":["How well did the team articulate the problem statement and support it with relevant statistics, data or evidence?","Is the proposed solution addressing a meaningful need for the identified users/stakeholders?","How well has the team researched existing solutions and competitors?","Is the proposed solution feasible and unique?","How would you rate the overall quality of the Ideation submission?","Additional feedback to candidate"],
"Design":["How effectively does the design demonstrate the proposed solution?","How well does the design address the identified problem?","How effectively has the team demonstrated the core features?","How feasible is the design from a technical perspective?","How would you rate the overall quality of the Design submission?","Additional feedback to candidate"],
"Build":["How effectively has the team implemented the proposed solution?","How well does the implementation address the identified problem?","How effectively has the team demonstrated the core features?","How feasible is the implementation from a technical perspective?","How would you rate the overall quality of the Build submission?","Additional feedback to candidate"],
"Demo":["How effectively does the final solution address the stated problem?","How well does the solution demonstrate its key features?","How effectively has the team demonstrated testing and validation?","How feasible and sustainable is the proposed solution?","How would you rate the overall quality of the final submission?","Additional feedback to candidate"]}

def extract_document(filename,data):
 ext=Path(filename or '').suffix.lower()
 if ext in ['.txt','.md','.csv']: text=data.decode('utf-8-sig',errors='replace')
 elif ext=='.pdf':
  with fitz.open(stream=data,filetype='pdf') as pdf: text='\n'.join(p.get_text() for p in pdf)
 elif ext=='.docx':
  d=docx.Document(io.BytesIO(data)); parts=[p.text for p in d.paragraphs]
  for t in d.tables:
   for row in t.rows: parts.append(' | '.join(c.text for c in row.cells))
  text='\n'.join(parts)
 elif ext=='.pptx':
  prs=Presentation(io.BytesIO(data)); slides=[]
  for n,s in enumerate(prs.slides,1):
   parts=[f'Slide {n}']
   for sh in s.shapes:
    if getattr(sh,'has_text_frame',False) and sh.text.strip(): parts.append(sh.text.strip())
    if getattr(sh,'has_table',False):
     for row in sh.table.rows: parts.append(' | '.join(c.text for c in row.cells))
   try:
    if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip(): parts.append('Speaker notes: '+s.notes_slide.notes_text_frame.text)
   except Exception: pass
   slides.append('\n'.join(parts))
  text='\n\n'.join(slides)
 elif ext=='.rtf': text=rtf_to_text(data.decode('utf-8',errors='replace'))
 else: raise ValueError('Supported files: TXT, PDF, DOCX, PPTX, RTF, MD, CSV. Convert legacy DOC/PPT files first.')
 if not text.strip(): raise ValueError('No readable text found. Scanned PDFs or image-only slides require OCR.')
 return text.strip()[:50000]

def parse_json(s):
 s=re.sub(r'^```(?:json)?\s*|\s*```$','',(s or '').strip(),flags=re.I)
 try: v=json.loads(s)
 except json.JSONDecodeError:
  m=re.search(r'\{.*\}',s,re.S)
  if not m: raise ValueError('Gemini returned invalid JSON.')
  v=json.loads(m.group())
 if not isinstance(v,dict): raise ValueError('Gemini response must be a JSON object.')
 return v

def normalize_rating(v):
 v=str(v or '').strip().lower()
 for phrase,rating in [('needs improvement','Needs Improvement'),('outstanding','Outstanding'),('very good','Very Good'),('incomplete','Incomplete'),('good','Good')]:
  if re.search(r'\b'+re.escape(phrase)+r'\b',v): return rating
 return 'Incomplete'

def normalize_answers(raw,questions):
 lookup={}
 for x in raw if isinstance(raw,list) else []:
  if isinstance(x,dict) and str(x.get('question','')).strip(): lookup[str(x['question']).strip()]=x.get('answer',x.get('rating',''))
 out=[]
 for i,q in enumerate(questions):
  a=lookup.get(q,'')
  a=normalize_rating(a) if i<5 else (str(a or '').strip() or 'No additional feedback provided.')
  out.append({'question':q,'answer':a})
 return out

def normalize_compliance(raw):
 if not isinstance(raw,list): return []
 valid={'satisfied':'Satisfied','partially satisfied':'Partially Satisfied','not satisfied':'Not Satisfied','not applicable':'Not Applicable'}; out=[]
 for x in raw:
  if not isinstance(x,dict): continue
  req=str(x.get('requirement','')).strip()
  if req: out.append({'requirement':req,'status':valid.get(str(x.get('status','')).strip().lower(),'Not Satisfied'),'evidence':str(x.get('evidence','')).strip() or 'No supporting evidence identified.','gap':str(x.get('gap','')).strip() or 'No additional gap explanation provided.'})
 return out

def create_prompt(name, project_id, project_name, phase, guidelines="", submission="", sample_reference=""):
 qs = PHASES[phase]
 qtext = "\n".join(f"Q{i}: {q}" for i, q in enumerate(qs, 1))
 if phase == "Ideation":
  guideline_section = guidelines[:20000] if guidelines.strip() else "No readable guideline text was supplied."
  sample_section = sample_reference[:15000] if sample_reference.strip() else "No separate sample project PPT was supplied. Do not infer mandatory requirements from a missing sample."
  ideation_instructions = '''
PHASE SCOPE — IDEATION ONLY:
The guideline document and sample project PPT apply ONLY to the Ideation phase.
The OFFICIAL GUIDELINES define explicit mandatory requirements and evaluation criteria. The SAMPLE PROJECT PPT is only an example to help students understand structure, presentation, and expected level of detail. Its topic, domain, project title, technology, features, and solution are NOT mandatory just because they appear in the sample.
Evaluate the student's own project idea. Do not penalize a different topic from the sample. Extract requirements for Part B only from explicit official guideline requirements, not illustrative sample content. If formal requirements and examples appear mixed together and cannot be distinguished, do not assume the example topic is mandatory.
'''
  compliance_instructions = '''PART B — IDEATION GUIDELINE COMPLIANCE:
Extract every distinct actionable requirement explicitly stated in the official Ideation guidelines. Do not convert the sample PPT's topic/title/features into requirements. Search the student submission for semantic evidence; synonyms and equivalent wording count. For each requirement return status exactly one of Satisfied, Partially Satisfied, Not Satisfied, Not Applicable. Missing evidence is Not Satisfied, not Not Applicable. Include requirement, evidence (quote briefly or faithfully paraphrase; say when none is found), and gap/improvement. A feature being mentioned is not proof it was implemented or tested. Do not invent statistics, tests, page numbers, or citations. If no explicit requirements can be extracted, return an empty list and mention the limitation in feedback. Keep compliance separate from phase ratings.'''
 else:
  guideline_section = "NOT APPLICABLE: This guideline document is used only for the Ideation phase. Do not use or infer any requirements from it for this evaluation."
  sample_section = "NOT APPLICABLE: The sample project PPT is used only for the Ideation phase. Do not use its topic, domain, content, format, or features to evaluate this phase."
  ideation_instructions = '''
PHASE SCOPE — NON-IDEATION:
This is a Design, Build, or Demo evaluation. Evaluate the student's actual submission using only the selected phase's questions and the phase-appropriate evidence present in the student submission. The Ideation guideline document and sample project PPT do not apply to this phase and must not affect ratings, feedback, suggestions, or requirements. Do not include any claim that the student failed to follow the Ideation guideline/sample PPT.
'''
  compliance_instructions = '''PART B — GUIDELINE COMPLIANCE:
Guideline compliance is only performed during the Ideation phase. For this Design, Build, or Demo evaluation, return an empty guideline_compliance array: []. Do not extract or assess requirements from the Ideation guideline document or sample project PPT.'''
 return f'''You are Evalora AI, an evidence-based project evaluator. Perform (A) phase-specific evaluation and (B) guideline compliance only where applicable. Compare meaning and evidence, not literal wording. Never invent evidence.

STUDENT DETAILS: Name={name}; Project ID={project_id}; Project={project_name}; Phase={phase}
PHASE QUESTIONS:
{qtext}
{ideation_instructions}
OFFICIAL GUIDELINES INPUT:
{guideline_section}

SAMPLE PROJECT PPT / REFERENCE INPUT:
{sample_section}

STUDENT SUBMISSION:
{submission[:30000]}

PART A — PHASE-SPECIFIC RATINGS:
Evaluate Q1-Q5 independently using the student's submission and the selected phase's criteria. For Ideation, use explicit official guidelines as relevant but never treat sample topic/content as mandatory. For Design, Build, and Demo, use the student submission and the selected phase's questions only; do not use the Ideation guideline or sample PPT. Return only these exact ratings: Outstanding, Very Good, Good, Needs Improvement, Incomplete. Outstanding=strong evidence/no major gaps; Very Good=minor gaps; Good=some meaningful details missing; Needs Improvement=major gaps; Incomplete=absent/not assessable. Q5 is an independent overall quality rating. Q6 is constructive feedback, not a rating. Give specific actionable suggestions grounded in the student's own project. ai_confidence_score is 0-100 confidence in the readability and completeness of the student evidence, not confidence that the project matches the sample. answer_accuracy is a separate AI-estimated 0-100 assessment of how accurately the ratings, feedback, and suggestions reflect evidence in the student's submission. Be conservative: deduct for unsupported claims, overlooked contradictory evidence, or conclusions not justified by the submission. This is an AI estimate, not independently verified ground-truth accuracy. Do not return overall score; app calculates it. Mapping: Outstanding 100, Very Good 85, Good 70, Needs Improvement 40, Incomplete 0.

{compliance_instructions}

Return valid JSON only, no markdown fences, with exactly this structure and actual results:
{{"phase_answers":[{{"question":"{qs[0]}","answer":"Good"}},{{"question":"{qs[1]}","answer":"Very Good"}},{{"question":"{qs[2]}","answer":"Needs Improvement"}},{{"question":"{qs[3]}","answer":"Good"}},{{"question":"{qs[4]}","answer":"Needs Improvement"}},{{"question":"{qs[5]}","answer":"Constructive feedback."}}],"ai_confidence_score":80,"answer_accuracy":90,"feedback":"Evidence-based overall feedback.","suggestions":"Specific actionable suggestions.","guideline_compliance":[]}}
Example values are structural only; generate actual answers from the documents. For Ideation, populate guideline_compliance with actual requirement-level results; for Design, Build, or Demo, return an empty array.'''

def call_gemini(prompt):
 key=os.getenv('GEMINI_API_KEY')
 if not key: raise ValueError('GEMINI_API_KEY is missing from Vercel Environment Variables.')
 client=genai.Client(api_key=key); errors=[]
 for model in MODELS:
  try:
   res=client.models.generate_content(model=model,contents=prompt,config=types.GenerateContentConfig(response_mime_type='application/json',temperature=0.2))
   if not res.text: raise ValueError('Gemini returned an empty response.')
   return parse_json(res.text),model
  except Exception as e: errors.append(f'{model}: {str(e)[:220]}')
 raise RuntimeError('All configured Gemini models failed. Verify model IDs and API access.\n'+'\n'.join(errors))

def connect_sheets():
 raw=os.getenv('GOOGLE_SERVICE_ACCOUNT_JSON')
 if not raw: raise ValueError('GOOGLE_SERVICE_ACCOUNT_JSON is missing from Vercel Environment Variables.')
 creds=Credentials.from_service_account_info(json.loads(raw),scopes=['https://www.googleapis.com/auth/spreadsheets','https://www.googleapis.com/auth/drive'])
 book=gspread.authorize(creds).open_by_key(SPREADSHEET_ID)
 try: ws=book.worksheet(WORKSHEET)
 except gspread.WorksheetNotFound: ws=book.add_worksheet(title=WORKSHEET,rows=1000,cols=len(HEADERS))
 if ws.col_count < len(HEADERS): ws.resize(cols=len(HEADERS))
 headers=ws.row_values(1)
 if not headers:
  ws.update(range_name='A1:K1',values=[HEADERS],value_input_option='RAW')
 else:
  if headers[:len(BASE_HEADERS)] != BASE_HEADERS:
   raise ValueError('Worksheet row 1 must begin with these nine headers: '+' | '.join(BASE_HEADERS))
  # Add the two requested columns after the existing nine columns, preserving old records.
  for offset, expected in enumerate(EXTRA_HEADERS, start=len(BASE_HEADERS)):
   existing = headers[offset] if len(headers) > offset else ''
   if existing and existing != expected:
    raise ValueError(f'Google Sheets column {offset + 1} must be named {expected!r}, not {existing!r}.')
  ws.update(range_name='J1:K1',values=[EXTRA_HEADERS],value_input_option='RAW')
 return ws

def append_to_google_sheet(row):
 if len(row)!=len(HEADERS): raise ValueError(f'Google Sheets expects exactly {len(HEADERS)} columns.')
 connect_sheets().append_row(row,value_input_option='RAW')

def compliance_report(items, phase="Ideation"):
 lines=['### Guideline vs Student Submission','','Assessment compares meaning and evidence, not literal text matching.','']
 if phase != "Ideation":
  return '\n'.join(lines+['**Not evaluated for this phase. The guideline document and sample project PPT are used only for Ideation.**'])
 if not items: return '\n'.join(lines+['**Compliance assessment unavailable.** No usable requirement-level results were returned.'])
 labels={'Satisfied':'SATISFIED','Partially Satisfied':'PARTIALLY SATISFIED','Not Satisfied':'NOT SATISFIED','Not Applicable':'NOT APPLICABLE'}; counts={k:0 for k in labels}
 for i,x in enumerate(items,1):
  s=x['status']; counts[s]+=1; lines += [f"#### Requirement {i}: {x['requirement']}",'',f"**Status:** {labels[s]}",'',f"**Student Evidence:** {x['evidence']}",'',f"**Gap / Improvement:** {x['gap']}",'','---','']
 applicable=len(items)-counts['Not Applicable']; lines += ['### Compliance Summary','',f'- Total requirements identified: {len(items)}',f"- Satisfied: {counts['Satisfied']}",f"- Partially Satisfied: {counts['Partially Satisfied']}",f"- Not Satisfied: {counts['Not Satisfied']}",f"- Not Applicable: {counts['Not Applicable']}"]
 lines += [f"- Fully satisfied applicable requirements: {round(counts['Satisfied']/applicable*100,2)}%" if applicable else '- No applicable requirements available for percentage calculation.','', 'This percentage is not the overall project score.']
 return '\n'.join(lines)

@app.get('/api/health')
def health(): return {'ok':True,'service':'Evalora AI','models':MODELS}
@app.get('/api/phases')
def phases(): return {'phases':list(PHASES),'questions':PHASES}
@app.post('/api/evaluate')
async def evaluate(
 name: str = Form(...),
 project_id: str = Form(...),
 project_name: str = Form(...),
 phase: str = Form(...),
 guideline_file: UploadFile | None = File(None),
 submission_file: UploadFile = File(...),
 sample_file: UploadFile | None = File(None),
):
 if phase not in PHASES:
  raise HTTPException(400, 'Select a valid phase.')
 try:
  sd = await submission_file.read()
  if len(sd) > 8 * 1024 * 1024:
   raise HTTPException(413, 'The student submission must be 8 MB or smaller.')
  submission = extract_document(submission_file.filename, sd)

  # Guidelines and sample reference are used in Ideation only.
  guidelines = ""
  sample_text = ""
  if phase == "Ideation":
   if guideline_file is None or not guideline_file.filename:
    raise HTTPException(400, 'Upload the official evaluation guideline for the Ideation phase.')
   gd = await guideline_file.read()
   if len(gd) > 8 * 1024 * 1024:
    raise HTTPException(413, 'The Ideation guideline file must be 8 MB or smaller.')
   guidelines = extract_document(guideline_file.filename, gd)
   if sample_file is not None and sample_file.filename:
    smd = await sample_file.read()
    if len(smd) > 8 * 1024 * 1024:
     raise HTTPException(413, 'The sample project PPT must be 8 MB or smaller.')
    sample_text = extract_document(sample_file.filename, smd)
  else:
   # Intentionally do not read, extract, or send the Ideation guideline/sample content
   # to Gemini for Design, Build, or Demo evaluations.
   guidelines = ""
   sample_text = ""

  result, model = call_gemini(create_prompt(name, project_id, project_name, phase, guidelines, submission, sample_text))
  answers = normalize_answers(result.get('phase_answers', []), PHASES[phase])
  compliance = normalize_compliance(result.get('guideline_compliance', [])) if phase == 'Ideation' else []
  score = round(sum(RATING_SCORES[x['answer']] for x in answers[:5]) / 5, 2)
  try:
   confidence = float(result.get('ai_confidence_score', 0))
  except (ValueError, TypeError):
   confidence = 0
  confidence = max(0, min(100, confidence))
  try:
   answer_accuracy = round(max(0, min(100, float(result.get('answer_accuracy')))), 2)
  except (ValueError, TypeError):
   answer_accuracy = ''
  feedback = str(result.get('feedback', '')).strip() or 'No overall feedback provided.'
  suggestions = str(result.get('suggestions', '')).strip() or 'No suggestions provided.'
  questions = '\n\n'.join((f"Q{i}: {x['question']}\nRating: {x['answer']}" if i <= 5 else f"Q{i}: {x['question']}\nFeedback: {x['answer']}") for i, x in enumerate(answers, 1))
  append_to_google_sheet([name, project_id, project_name, phase, questions, confidence, feedback, suggestions, score, model, answer_accuracy])
  lines = ['# Evaluation Results', '', f'**Student:** {name}', f'**Project ID:** {project_id}', f'**Project:** {project_name}', f'**Phase:** {phase}', '', '## Overall Score: ' + str(score) + '/100', '', f'**AI Confidence:** {confidence}/100', '', '## Phase Ratings', '']
  for i, x in enumerate(answers[:5], 1):
   lines += [f"### Q{i}: {x['question']}", '', f"**Rating:** {x['answer']}", '']
  lines += ['## Additional Feedback', '', answers[5]['answer'], '', '## Overall Feedback', '', feedback, '', '## Suggestions', '', suggestions, '', compliance_report(compliance, phase)]
  compliance_status = 'Ideation guideline compliance included.' if phase == 'Ideation' else 'Ideation guideline/sample PPT skipped for this phase.'
  return {'success': True, 'report': '\n'.join(lines), 'status': f'Evaluation completed using {model}. Saved to Google Sheets. {compliance_status}', 'overall_score': score, 'confidence': confidence, 'model_used': model, 'sheet_saved': True}
 except HTTPException:
  raise
 except Exception as e:
  raise HTTPException(500, str(e)[:1200])
