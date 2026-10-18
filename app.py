import os

import re

import json

import time

import shutil

import tempfile

import subprocess



import gradio as gr

import pandas as pd

import pymupdf

import docx

import gspread



from pptx import Presentation

from striprtf.striprtf import rtf_to_text

from dotenv import load_dotenv

from google import genai

from google.genai import types

from google.oauth2.service_account import Credentials





# =====================================================

# CONFIGURATION

# =====================================================



load_dotenv()



GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

SERVICE_ACCOUNT_FILE = "service_account.json"



SPREADSHEET_ID = "1L8FuY_CFURGrRkWbEFAG_QsHTyF0oUJYfpCfNAl-sOw"

WORKSHEET_NAME = "Evaluations"



# Keep this list unchanged.

MODELS = [

    "gemini-3.8-flash",

    "gemini-3.6-flash",

    "gemini-3.5-flash-lite",

]



RATINGS = [

    "Outstanding",

    "Very Good",

    "Good",

    "Needs Improvement",

    "Incomplete",

]



RATING_SCORES = {

    "Outstanding": 100,

    "Very Good": 85,

    "Good": 70,

    "Needs Improvement": 40,

    "Incomplete": 0,

}



SHEET_HEADERS = [

    "Name",

    "Project ID",

    "Project Name",

    "Phase Type",

    "Questions",

    "AI Confidence Score",

    "Feedback",

    "Suggestions",

    "Overall Score",

]



SUPPORTED_FILES = [

    ".txt", ".pdf", ".doc", ".docx",

    ".ppt", ".pptx", ".rtf",

]





# =====================================================

# EVALUATION QUESTIONS

# =====================================================



PHASES = {

    "Ideation": [

        "How well did the team articulate the problem statement and support it with relevant statistics, data or evidence?",

        "Is the proposed solution addressing a meaningful need for the identified users/stakeholders?",

        "How well has the team researched existing solutions and competitors?",

        "Is the proposed solution feasible and unique?",

        "How would you rate the overall quality of the Ideation submission?",

        "Additional feedback to candidate",

    ],

    "Design": [

        "How effectively does the design demonstrate the proposed solution?",

        "How well does the design address the identified problem?",

        "How effectively has the team demonstrated the core features?",

        "How feasible is the design from a technical perspective?",

        "How would you rate the overall quality of the Design submission?",

        "Additional feedback to candidate",

    ],

    "Build": [

        "How effectively has the team validated the problem with target users?",

        "Is there evidence that the solution meets user needs?",

        "How effectively has the team collected and analyzed user feedback?",

        "How well has the team demonstrated feasibility based on Build?",

        "How would you rate the overall quality of the Build submission?",

        "Additional feedback to candidate",

    ],

    "Demo": [

        "How effectively does the final solution address the stated problem?",

        "How well does the solution demonstrate its key features?",

        "How effectively has the team demonstrated testing and Build?",

        "How feasible and sustainable is the proposed solution?",

        "How would you rate the overall quality of the final submission?",

        "Additional feedback to candidate",

    ],

}





# =====================================================

# DOCUMENT EXTRACTION

# TXT, PDF, DOC, DOCX, PPT, PPTX, RTF

# =====================================================



def extract_document(path):

    if not path:

        raise ValueError("Please upload a document.")



    extension = os.path.splitext(path)[1].lower()

    text = ""



    if extension == ".txt":

        for encoding in ("utf-8-sig", "utf-8", "cp1252"):

            try:

                with open(path, "r", encoding=encoding) as file:

                    text = file.read()

                break

            except UnicodeDecodeError:

                continue



    elif extension == ".pdf":

        with pymupdf.open(path) as pdf:

            text = "\n".join(page.get_text() for page in pdf)



    elif extension == ".docx":

        document = docx.Document(path)

        parts = [p.text for p in document.paragraphs]



        for table in document.tables:

            for row in table.rows:

                parts.append(

                    " | ".join(cell.text for cell in row.cells)

                )



        text = "\n".join(parts)



    elif extension == ".pptx":

        presentation = Presentation(path)

        slides = []



        for number, slide in enumerate(presentation.slides, start=1):

            parts = [f"Slide {number}"]



            for shape in slide.shapes:

                if getattr(shape, "has_text_frame", False):

                    content = shape.text.strip()

                    if content:

                        parts.append(content)



                if getattr(shape, "has_table", False):

                    for row in shape.table.rows:

                        parts.append(

                            " | ".join(cell.text for cell in row.cells)

                        )



            slides.append("\n".join(parts))



        text = "\n\n".join(slides)



    elif extension == ".rtf":

        with open(path, "r", encoding="utf-8", errors="replace") as file:

            text = rtf_to_text(file.read())



    elif extension in (".doc", ".ppt"):

        # Legacy Office files are converted using LibreOffice.

        if not shutil.which("libreoffice"):

            raise RuntimeError(

                f"{extension} files require LibreOffice. Install it and "

                "ensure the 'libreoffice' command is available, or save "

                "the file as DOCX/PPTX."

            )



        output_dir = tempfile.mkdtemp(prefix="evalora_")

        target_extension = ".docx" if extension == ".doc" else ".pptx"

        target_format = "docx" if extension == ".doc" else "pptx"



        try:

            result = subprocess.run(

                [

                    "libreoffice",

                    "--headless",

                    "--convert-to", target_format,

                    "--outdir", output_dir,

                    path,

                ],

                capture_output=True,

                text=True,

                timeout=90,

            )



            target_name = (

                os.path.splitext(os.path.basename(path))[0]

                + target_extension

            )

            target_path = os.path.join(output_dir, target_name)



            if result.returncode != 0 or not os.path.exists(target_path):

                raise RuntimeError(

                    "Office file conversion failed: "

                    + (result.stderr or result.stdout)

                )



            text = extract_document(target_path)



        finally:

            shutil.rmtree(output_dir, ignore_errors=True)



    else:

        raise ValueError(

            "Supported formats: TXT, PDF, DOC, DOCX, PPT, PPTX, RTF."

        )



    if not text or not text.strip():

        raise ValueError(

            "No readable text found. Scanned PDFs or image-only files "

            "need OCR before evaluation."

        )



    return text.strip()





# =====================================================

# JSON PARSER

# =====================================================



def parse_json(response_text):

    cleaned = response_text.strip()



    cleaned = re.sub(

        r"^```(?:json)?\s*",

        "",

        cleaned,

        flags=re.IGNORECASE,

    )

    cleaned = re.sub(r"\s*```$", "", cleaned)



    try:

        data = json.loads(cleaned)

    except json.JSONDecodeError:

        match = re.search(r"\{.*\}", cleaned, re.DOTALL)



        if not match:

            raise ValueError("Gemini returned invalid JSON.")



        data = json.loads(match.group(0))



    if not isinstance(data, dict):

        raise ValueError("Gemini response must be a JSON object.")



    return data





# =====================================================

# RATING VALIDATION

# =====================================================



def normalize_rating(value):

    text = str(value or "").strip().lower()



    # Match longer phrases before shorter phrases.

    rating_lookup = [

        ("needs improvement", "Needs Improvement"),

        ("outstanding", "Outstanding"),

        ("very good", "Very Good"),

        ("incomplete", "Incomplete"),

        ("good", "Good"),

    ]



    for phrase, rating in rating_lookup:

        if re.search(r"\b" + re.escape(phrase) + r"\b", text):

            return rating



    return "Incomplete"





def normalize_answers(raw_answers, questions):

    lookup = {}



    if isinstance(raw_answers, list):

        for item in raw_answers:

            if not isinstance(item, dict):

                continue



            question = str(item.get("question", "")).strip()

            answer = item.get("answer", item.get("rating", ""))



            if question:

                lookup[question] = answer



    answers = []



    for index, question in enumerate(questions):

        raw_answer = lookup.get(question, "")



        if index < 5:

            answer = normalize_rating(raw_answer)

        else:

            answer = str(raw_answer or "").strip()

            if not answer:

                answer = "No additional feedback provided."



        answers.append({

            "question": question,

            "answer": answer,

        })



    return answers





# =====================================================

# GEMINI PROMPT

# =====================================================



def create_prompt(name, project_id, project_name, phase,
                  guidelines, submission):
    questions = PHASES[phase]
    question_text = "\\n".join(
        f"Q{i}: {question}"
        for i, question in enumerate(questions, 1)
    )

    return f"""
You are Evalora AI, a consistent, evidence-based project evaluator.
Evaluate the STUDENT SUBMISSION against the EVALUATION GUIDELINES.
The guidelines are criteria/reference material, not a second student project.
Do not score or describe the guideline document as if it were the submission.

STUDENT
Name: {name}
Project ID: {project_id}
Project Name: {project_name}
Phase: {phase}

QUESTIONS
{question_text}

EVALUATION GUIDELINES / REFERENCE
{guidelines[:25000]}

STUDENT SUBMISSION
{submission[:40000]}

MANDATORY EVALUATION PROCESS
1. Evaluate each Q1-Q5 separately against the exact criterion in the question.
2. Base ratings only on evidence actually present in the submission and relevant
   guideline requirements. Do not award marks merely for polished slides, long text,
   technical buzzwords, or a plausible idea.
3. Do not require a feature that the guidelines/question do not require. Do not
   penalize a submission for information not expected at its current phase.
4. Do not invent statistics, user research, competitors, tests, citations, features,
   or implementation details. If evidence is absent, treat it as absent.
5. Apply these rating anchors to EACH criterion, using criterion-specific evidence:
   Outstanding (100): Strong, specific, relevant evidence fully satisfies the criterion;
   important claims are substantiated and there are no major gaps.
   Very Good (85): Addressed well with relevant evidence; only minor gaps remain.
   Good (70): Reasonably addressed, but meaningful evidence/detail is missing or claims
   are partly generic.
   Needs Improvement (40): Only partly addressed; mostly generic, weakly supported,
   overly broad, or has major evidence gaps.
   Incomplete (0): Absent, not assessable from the submission, or fundamentally fails
   to address the criterion.
6. Interpret Q1-Q4 according to their exact wording. Do not rate problem evidence highly
   without relevant data; do not rate competitor research highly unless named existing
   solutions/competitors are compared; do not rate feasibility/uniqueness highly based
   only on a generic workflow or list of tools.
7. Q5 is an OVERALL QUALITY rating, not a rating of visual presentation alone. Consider
   completeness, coherence, phase-appropriate evidence, alignment with the problem,
   and major unresolved gaps. Use the same anchors. Judge Q5 independently.
8. Q1-Q5 answers must contain exactly one allowed rating and nothing else:
   Outstanding, Very Good, Good, Needs Improvement, Incomplete.
9. Q6 must contain concise, constructive additional feedback. If no further feedback is
   justified, say so explicitly.
10. Feedback must explain the strongest evidence and most important gaps. Suggestions
    must be actionable and tied to the gaps found.
11. ai_confidence_score is confidence that the supplied documents contain enough
    readable evidence to support this evaluation (0-100), NOT project quality and NOT
    statistical accuracy. Lower it for incomplete, ambiguous, or hard-to-read files.
12. Do not return an overall score. The application calculates it deterministically
    from Q1-Q5: Outstanding=100, Very Good=85, Good=70,
    Needs Improvement=40, Incomplete=0. The arithmetic average is the final score.

Return valid JSON only, with no Markdown fences, using this exact shape:
{{
  "phase_answers": [
    {{"question": "{questions[0]}", "answer": "Good"}},
    {{"question": "{questions[1]}", "answer": "Very Good"}},
    {{"question": "{questions[2]}", "answer": "Needs Improvement"}},
    {{"question": "{questions[3]}", "answer": "Good"}},
    {{"question": "{questions[4]}", "answer": "Needs Improvement"}},
    {{"question": "{questions[5]}", "answer": "Constructive additional feedback."}}
  ],
  "ai_confidence_score": 80,
  "feedback": "Evidence-based overall assessment, including strengths and gaps.",
  "suggestions": "Specific actions that address the identified gaps."
}}
"""
# =====================================================

# GEMINI FALLBACK

# =====================================================



def call_gemini(prompt, progress=None):

    if not GEMINI_API_KEY:

        raise ValueError("GEMINI_API_KEY is missing from .env.")



    client = genai.Client(api_key=GEMINI_API_KEY)

    errors = []



    for index, model_name in enumerate(MODELS):

        try:

            print(f"Trying model: {model_name}")



            if progress:

                progress(

                    min(0.45 + index * 0.08, 0.68),

                    desc=f"Trying {model_name}...",

                )



            response = client.models.generate_content(

                model=model_name,

                contents=prompt,

                config=types.GenerateContentConfig(

                    response_mime_type="application/json",

                    temperature=0.2,

                ),

            )



            if not response.text:

                raise ValueError("Empty Gemini response.")



            result = parse_json(response.text)



            print(f"Success with model: {model_name}")

            return result, model_name



        except Exception as error:

            errors.append(f"{model_name}: {error}")

            print(f"Model failed: {model_name}: {error}")

            time.sleep(1)



    raise RuntimeError(

        "All configured Gemini models failed:\n"

        + "\n".join(errors)

    )





# =====================================================

# GOOGLE SHEETS

# =====================================================



def connect_google_sheets():

    if not os.path.exists(SERVICE_ACCOUNT_FILE):

        raise FileNotFoundError(

            "service_account.json was not found beside app.py."

        )



    scopes = [

        "https://www.googleapis.com/auth/spreadsheets",

        "https://www.googleapis.com/auth/drive",

    ]



    credentials = Credentials.from_service_account_file(

        SERVICE_ACCOUNT_FILE,

        scopes=scopes,

    )



    client = gspread.authorize(credentials)

    spreadsheet = client.open_by_key(SPREADSHEET_ID)



    try:

        worksheet = spreadsheet.worksheet(WORKSHEET_NAME)

    except gspread.WorksheetNotFound:

        worksheet = spreadsheet.add_worksheet(

            title=WORKSHEET_NAME,

            rows=1000,

            cols=9,

        )



    current_headers = worksheet.row_values(1)



    if not current_headers:

        worksheet.append_row(SHEET_HEADERS, value_input_option="RAW")

    elif current_headers != SHEET_HEADERS:

        raise ValueError(

            "Google Sheet row 1 must have these exact headers:\n"

            + " | ".join(SHEET_HEADERS)

        )



    return worksheet





def append_to_google_sheet(row_data):

    if len(row_data) != 9:

        raise ValueError("Expected exactly 9 Google Sheets columns.")



    worksheet = connect_google_sheets()

    worksheet.append_row(row_data, value_input_option="RAW")





# =====================================================

# EXCEL REPORT

# =====================================================



def create_excel_report(name, project_id, project_name, phase,

                        answers, confidence, feedback,

                        suggestions, overall_score):



    summary = pd.DataFrame([{

        "Name": name,

        "Project ID": project_id,

        "Project Name": project_name,

        "Phase Type": phase,

        "AI Confidence Score": confidence,

        "Feedback": feedback,

        "Suggestions": suggestions,

        "Overall Score": overall_score,

    }])



    question_rows = pd.DataFrame([

        {

            "Question Number": f"Q{i}",

            "Question": item["question"],

            "Answer": item["answer"],

        }

        for i, item in enumerate(answers, 1)

    ])



    fd, report_path = tempfile.mkstemp(suffix=".xlsx")

    os.close(fd)



    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:

        summary.to_excel(

            writer,

            sheet_name="Evaluation Summary",

            index=False,

        )

        question_rows.to_excel(

            writer,

            sheet_name="Question Ratings",

            index=False,

        )



    return report_path





# =====================================================

# EVALUATION

# =====================================================



def evaluate(name, project_id, project_name, phase,

             guideline_file, submission_file,

             progress=gr.Progress()):



    try:

        if not all([

            name, project_id, project_name, phase,

            guideline_file, submission_file

        ]):

            raise ValueError(

                "Fill in all project details and upload both documents."

            )



        if phase not in PHASES:

            raise ValueError("Select a valid phase.")



        progress(0.1, desc="Reading documents...")



        guidelines = extract_document(guideline_file)

        submission = extract_document(submission_file)



        progress(0.3, desc="Preparing evaluation...")



        prompt = create_prompt(

            name,

            project_id,

            project_name,

            phase,

            guidelines,

            submission,

        )



        result, used_model = call_gemini(

            prompt,

            progress=progress,

        )



        progress(0.72, desc="Validating ratings...")



        answers = normalize_answers(

            result.get("phase_answers", []),

            PHASES[phase],

        )



        overall_score = round(

            sum(

                RATING_SCORES[item["answer"]]

                for item in answers[:5]

            ) / 5,

            2,

        )



        try:

            confidence = float(result.get("ai_confidence_score", 0))

        except (ValueError, TypeError):

            confidence = 0



        confidence = max(0, min(100, confidence))



        feedback = str(

            result.get("feedback", "")

        ).strip() or "No overall feedback provided."



        suggestions = str(

            result.get("suggestions", "")

        ).strip() or "No suggestions provided."



        # Store all six questions and answers in ONE cell.

        questions_text = "\n\n".join(

            f"Q{i}: {item['question']}\nAnswer: {item['answer']}"

            for i, item in enumerate(answers, 1)

        )



        row_data = [

            name,

            project_id,

            project_name,

            phase,

            questions_text,

            confidence,

            feedback,

            suggestions,

            overall_score,

        ]



        progress(0.8, desc="Saving to Google Sheets...")

        append_to_google_sheet(row_data)



        progress(0.9, desc="Generating Excel report...")



        report_path = create_excel_report(

            name,

            project_id,

            project_name,

            phase,

            answers,

            confidence,

            feedback,

            suggestions,

            overall_score,

        )



        ratings_df = pd.DataFrame([

            {

                "Question": item["question"],

                "Answer": item["answer"],

            }

            for item in answers

        ])



        summary_df = pd.DataFrame([

            {"Field": "Model Used", "Value": used_model},

            {"Field": "AI Confidence Score", "Value": confidence},

            {"Field": "Feedback", "Value": feedback},

            {"Field": "Suggestions", "Value": suggestions},

            {"Field": "Overall Score", "Value": overall_score},

        ])



        rating_text = "\n".join(

            f"- **Q{i}:** {item['answer']}"

            for i, item in enumerate(answers[:5], 1)

        )



        report = f"""

# Evalora AI Evaluation Report



**Student:** {name}



**Project ID:** {project_id}



**Project Name:** {project_name}



**Phase:** {phase}



**Gemini Model:** `{used_model}`



**Overall Score:** {overall_score}/100



**AI Confidence Score:** {confidence}/100



## Question Ratings



{rating_text}



## Additional Feedback (Q6)



{answers[5]["answer"]}



## Overall Feedback



{feedback}



## Suggestions



{suggestions}

"""



        progress(1.0, desc="Evaluation completed.")



        status = (

            f"Evaluation completed using {used_model}. "

            "Saved to Google Sheets and Excel report generated."

        )



        return (

            report,

            ratings_df,

            summary_df,

            report_path,

            status,

        )



    except Exception as error:

        message = f"Evaluation failed: {error}"



        return (

            f"## Error\n\n{message}",

            pd.DataFrame(columns=["Question", "Answer"]),

            pd.DataFrame(columns=["Field", "Value"]),

            None,

            message,

        )





# =====================================================

# GRADIO WEBPAGE

# =====================================================



with gr.Blocks(

    title="Evalora AI",

    theme=gr.themes.Soft(),

) as demo:



    gr.Markdown(

        """

        # Evalora AI

        ### AI-powered project evaluation



        **Supported:** TXT, PDF, DOC, DOCX, PPT, PPTX, RTF



        Q1-Q5 return ratings only. Q6 provides additional feedback.

        """

    )



    with gr.Row():

        name_input = gr.Textbox(label="Student Name")

        project_id_input = gr.Textbox(label="Project ID")



    project_name_input = gr.Textbox(label="Project Name")



    phase_input = gr.Dropdown(

        choices=list(PHASES.keys()),

        value="Ideation",

        label="Phase Type",

    )



    with gr.Row():

        guideline_input = gr.File(

            label="Evaluation Guidelines",

            file_types=SUPPORTED_FILES,

            type="filepath",

        )



        submission_input = gr.File(

            label="Student Submission",

            file_types=SUPPORTED_FILES,

            type="filepath",

        )



    evaluate_button = gr.Button(

        "Evaluate Project",

        variant="primary",

    )



    status_output = gr.Markdown()

    report_output = gr.Markdown()



    with gr.Accordion("Question Ratings", open=True):

        ratings_output = gr.Dataframe(

            label="Q1-Q6",

            interactive=False,

        )



    with gr.Accordion("Evaluation Summary", open=True):

        summary_output = gr.Dataframe(

            label="Evaluation Summary",

            interactive=False,

        )



    excel_output = gr.File(label="Download Excel Report")



    evaluate_button.click(

        fn=evaluate,

        inputs=[

            name_input,

            project_id_input,

            project_name_input,

            phase_input,

            guideline_input,

            submission_input,

        ],

        outputs=[

            report_output,

            ratings_output,

            summary_output,

            excel_output,

            status_output,

        ],

        api_name="evaluate",

    )





# =====================================================

# START SERVER

# =====================================================



if __name__ == "__main__":

    demo.queue().launch(
        server_name="0.0.0.0",
        server_port=int(os.environ.get("PORT", "7860")),
        share=False,
        debug=False,
    )
